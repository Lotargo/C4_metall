"""Read the supplied attachments without changing them; build website previews."""
from pathlib import Path
from hashlib import sha256
import json
from datetime import date, datetime
import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'public/documents/implementation'
assets = {}


def cell_data(cell):
    value = cell.value
    if value is None:
        return {'value': None, 'text': ''}
    if cell.data_type == 'f':
        raise ValueError(f'Formula needs an explicit display policy: {cell.coordinate}')
    if isinstance(value, (date, datetime)):
        text = value.isoformat()
        value = text
    elif isinstance(value, bool):
        text = str(value)
    elif isinstance(value, (int, float)):
        if cell.number_format == '0.00%':
            text = f'{value * 100:.2f}%'.replace('.', ',')
        elif cell.number_format == '#,##0':
            text = f'{value:,.0f}'.replace(',', '\u202f')
        elif cell.number_format in ('General', '@'):
            text = format(value, '.15g').replace('.', ',')
        else:
            raise ValueError(f'Unsupported number format: {cell.number_format}')
    else:
        text = str(value)
    return {'value': value, 'text': text, 'format': cell.number_format}


for file in sorted(SOURCE.rglob('*')):
    if not file.is_file() or file.suffix not in ('.png', '.xlsx'):
        continue
    key = file.relative_to(SOURCE).as_posix()
    asset = {'sha256': sha256(file.read_bytes()).hexdigest()}
    if file.suffix == '.png':
        with Image.open(file) as image:
            asset.update(kind='image', width=image.width, height=image.height)
    elif file.suffix == '.xlsx':
        workbook = openpyxl.load_workbook(file, data_only=False)
        sheets = []
        for sheet in workbook:
            cells = [cell for row in sheet for cell in row if cell.value is not None]
            if not cells:
                sheets.append({'name': sheet.title, 'rows': [], 'range': None})
                continue
            if sheet.merged_cells.ranges:
                raise ValueError(f'Merged cells require explicit support: {sheet.title}')
            top, bottom = min(c.row for c in cells), max(c.row for c in cells)
            left, right = min(c.column for c in cells), max(c.column for c in cells)
            rows = [[cell_data(cell) for cell in row] for row in sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right)]
            sheets.append({'name': sheet.title, 'range': f'{sheet.cell(top, left).coordinate}:{sheet.cell(bottom, right).coordinate}', 'rows': rows})
        asset.update(kind='workbook', sheets=sheets)
    else:
        continue
    assets[key] = asset

output = ROOT / 'src/data/document-attachments.json'
output.write_text(json.dumps({'version': 1, 'assets': assets}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(f'Prepared {len(assets)} read-only attachment previews')
